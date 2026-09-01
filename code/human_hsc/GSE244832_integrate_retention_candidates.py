#!/usr/bin/env python3
"""Integrate donor HSC RNA, published HSC ATAC, and dynamic regression fates."""

import argparse
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
EDGE = REPO_ROOT / "results/GSE244832_hsc_donor"
DYNAMIC = REPO_ROOT / "results/GSE295293_dynamic_regression/gene_fates_HSC_Mesenchymal.tsv.gz"
SUPPLEMENT = REPO_ROOT / "data/raw_external/GSE244832/ScienceDirect_files_22Apr2025_06-13-10.681.zip"
OUT = EDGE / "integrated_retention_candidates.tsv"

FUNCTIONAL_SUPPORT = {
    "SERPINE1": "human spheroid knockdown; HSC-specific mouse knockout/pharmacology",
    "SPON1": "human HSC knockdown validation in source study",
    "LTBP2": "human HSC knockdown validation in source study",
    "EFEMP1": "human HSC knockdown validation in source study",
    "GAS7": "human HSC knockdown validation in source study",
    "RUNX1": "source-study promoter/motif and perturbation support",
    "RUNX2": "source-study promoter/motif and perturbation support",
}


def read_source_table() -> pd.DataFrame:
    with ZipFile(SUPPLEMENT) as archive:
        payload = archive.read("1-s2.0-S0168827824026679-mmc14.xlsx")
    sheet = load_workbook(BytesIO(payload), read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(min_row=8, values_only=True))
    header = [str(value) if value is not None else f"empty_{index}" for index, value in enumerate(rows[0])]
    result = pd.DataFrame(rows[1:], columns=header)
    result = result.rename(columns={result.columns[0]: "gene"})
    return result[result.gene.notna()].copy()


def read_edge(label: str) -> pd.DataFrame:
    data = pd.read_csv(EDGE / f"{label}_MASH_vs_nonactivated_edgeR.tsv", sep="\t")
    data["gene_upper"] = data.gene.str.upper()
    return data[["gene_upper", "logFC", "PValue", "FDR"]].rename(
        columns={column: f"{label}_{column}" for column in ("logFC", "PValue", "FDR")}
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--edge-dir", type=Path, default=EDGE)
    parser.add_argument("--dynamic-fates", type=Path, default=DYNAMIC)
    parser.add_argument("--source-supplement", type=Path, default=SUPPLEMENT)
    parser.add_argument("--output", type=Path, default=None)
    return parser.parse_args()


def main() -> None:
    global EDGE, DYNAMIC, SUPPLEMENT, OUT
    args = parse_args()
    EDGE, DYNAMIC, SUPPLEMENT = args.edge_dir, args.dynamic_fates, args.source_supplement
    OUT = args.output or EDGE / "integrated_retention_candidates.tsv"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    source = read_source_table()
    source["gene_upper"] = source.gene.astype(str).str.upper()
    source = source.rename(columns={
        "avg_log2FC_AvsQ_ATAC": "author_AvsQ_ATAC_logFC",
        "avg_log2FC_MASHvsNA_ATAC": "author_MASH_ATAC_logFC",
        "avg_log2FC_AvsQ_RNA": "author_AvsQ_RNA_logFC",
        "avg_log2FC_MASHvsNA_RNA": "author_MASH_RNA_logFC",
        "sum_sig": "author_sum_sig",
    })
    source = source[[
        "gene_upper", "author_AvsQ_ATAC_logFC", "author_MASH_ATAC_logFC",
        "author_AvsQ_RNA_logFC", "author_MASH_RNA_logFC", "author_sum_sig",
    ]].drop_duplicates("gene_upper")

    dynamic = pd.read_csv(DYNAMIC, sep="\t")
    dynamic["gene_upper"] = dynamic.gene.str.upper()
    dynamic = dynamic.drop(columns="gene").rename(columns={column: f"mouse_{column}" for column in dynamic.columns if column != "gene_upper"})

    merged = read_edge("metadata").merge(read_edge("supplement_swap"), on="gene_upper", how="outer")
    merged = merged.merge(source, on="gene_upper", how="left").merge(dynamic, on="gene_upper", how="left")
    merged["human_positive_both"] = (merged.metadata_logFC > 0) & (merged.supplement_swap_logFC > 0)
    merged["human_nominal_both"] = (merged.metadata_PValue < 0.05) & (merged.supplement_swap_PValue < 0.05)
    merged["human_fdr_any_010"] = np.minimum(merged.metadata_FDR, merged.supplement_swap_FDR) < 0.10
    merged["author_four_way"] = merged.author_sum_sig.eq(4)
    merged["mouse_positive_residual"] = (merged.mouse_residual_logFC > 0) & (merged.mouse_residual_FDR < 0.10)
    merged["mouse_persistent"] = merged.mouse_fate.eq("persistent")
    merged["cross_species_direction"] = merged.human_positive_both & (merged.mouse_progression_logFC > 0)
    merged["functional_support"] = merged.gene_upper.map(FUNCTIONAL_SUPPORT).fillna("")
    merged["evidence_score"] = (
        2 * merged.human_nominal_both.astype(int)
        + 2 * merged.mouse_persistent.astype(int)
        + merged.human_fdr_any_010.astype(int)
        + merged.author_four_way.astype(int)
        + merged.mouse_positive_residual.astype(int)
        + merged.cross_species_direction.astype(int)
        + merged.functional_support.ne("").astype(int)
    )
    merged["classification"] = np.select(
        [
            merged.human_positive_both & (merged.mouse_progression_logFC < 0),
            merged.mouse_persistent & merged.cross_species_direction & merged.author_four_way,
            merged.mouse_positive_residual & merged.human_positive_both,
        ],
        ["cross_species_direction_conflict", "pathogenic_retention_candidate", "residual_state_candidate"],
        default="insufficient_or_context_specific",
    )
    merged = merged.sort_values(["evidence_score", "metadata_PValue"], ascending=[False, True])
    merged.to_csv(OUT, sep="\t", index=False)
    columns = [
        "gene_upper", "evidence_score", "classification", "metadata_logFC", "metadata_FDR",
        "supplement_swap_logFC", "supplement_swap_FDR", "mouse_progression_logFC",
        "mouse_residual_logFC", "mouse_residual_FDR", "mouse_fate", "author_four_way",
        "functional_support",
    ]
    print(merged[columns].head(40).to_string(index=False))


if __name__ == "__main__":
    main()
